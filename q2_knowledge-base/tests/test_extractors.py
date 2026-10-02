from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.extractors import extract_directory, extract_file
from app.records import extract_metadata


def by_source(results):
    return {r["source"]: r for r in results}


def test_directory_statuses(raw_dir):
    results = by_source(extract_directory(raw_dir))

    assert ".DS_Store" not in results
    assert results["eligibility_criteria.md"]["status"] == "success"
    assert results["nested/partner_website.html"]["status"] == "success"
    assert results["repayment_policy.pdf"]["status"] == "success"
    assert results["scanned.pdf"]["status"] == "skipped"
    assert "scanned" in results["scanned.pdf"]["message"]
    assert results["corrupt.pdf"]["status"] == "error"
    assert results["notes.docx"]["status"] == "skipped"
    assert results["notes.docx"]["message"]  # skipped files always explain why


def test_html_keeps_sentences_and_tables_intact(raw_dir):
    doc = extract_file(raw_dir / "documents_required.html", root=raw_dir)

    assert doc["title"] == "Documents Required"
    assert "Home | Products" not in doc["text"]  # nav removed
    assert "Docs | Darwix" not in doc["text"]  # <head><title> is not body text
    assert "submit your GST returns and bank statements for the last" in doc["text"]
    assert "Document: PAN card | Mandatory: Yes" in doc["text"]


def test_pdf_text_and_page_count(raw_dir):
    doc = extract_file(raw_dir / "repayment_policy.pdf", root=raw_dir)

    assert doc["page_count"] == 1
    assert "Prepayment is allowed" in doc["text"]


def test_csv_rows_and_counts(raw_dir):
    doc = extract_file(raw_dir / "loan_limits_table.csv", root=raw_dir)

    assert doc["row_count"] == 2 and doc["column_count"] == 4
    assert "Loan Type: Term loan | Minimum Amount: 1000000" in doc["text"]


def test_csv_with_bom_and_semicolons(tmp_path):
    path = tmp_path / "t.csv"
    path.write_bytes("\ufeffName;Limit\nA;10\nB;20\n".encode("utf-8"))

    doc = extract_file(path)

    assert doc["text"].splitlines()[0] == "Name: A | Limit: 10"


def test_cp1252_file_is_decoded(tmp_path):
    path = tmp_path / "rupee.txt"
    path.write_bytes("Fee is \x80 500".encode("latin-1"))  # cp1252 euro sign

    assert "500" in extract_file(path)["text"]


def test_missing_directory_raises(tmp_path):
    import pytest

    with pytest.raises(FileNotFoundError):
        extract_directory(tmp_path / "nope")


# --- DOCX -----------------------------------------------------------------

W = nsdecls("w")


def docx_result(doc, tmp_path):
    path = tmp_path / "doc.docx"
    doc.save(path)

    return extract_file(path)


def fill(table, row, *texts, bold=False):
    for col, text in enumerate(texts):
        paragraph = table.cell(row, col).paragraphs[0]
        paragraph.add_run(text).bold = bold or None


def test_docx_headings_lists_and_table_in_document_order(tmp_path):
    doc = Document()
    doc.add_paragraph("Business Loan Guide", style="Title")
    doc.add_heading("Eligibility", level=1)
    doc.add_paragraph("Applicants need Rs 40 lakh turnover.")
    doc.add_paragraph("Three years in business", style="List Bullet")
    doc.add_paragraph("Credit score of 700", style="List Bullet")
    doc.add_heading("Fees", level=2)
    doc.add_paragraph("")  # empty paragraphs vanish
    table = doc.add_table(rows=3, cols=2)
    fill(table, 0, "Fee", "Amount", bold=True)
    fill(table, 1, "Processing", "2 percent")
    fill(table, 2, "Late payment", "Rs 500")
    doc.add_paragraph("After the table.")

    result = docx_result(doc, tmp_path)

    assert result["title"] == "Business Loan Guide"
    assert result["text"] == (
        "# Business Loan Guide\n# Eligibility\nApplicants need Rs 40 lakh turnover."
        "\n\n- Three years in business\n- Credit score of 700"
        "\n\n## Fees"
        "\nFee: Processing | Amount: 2 percent\nFee: Late payment | Amount: Rs 500"
        "\n\nAfter the table."
    )


def test_docx_table_layouts(tmp_path):
    doc = Document()

    # Header row flagged in Word (not bold).
    flagged = doc.add_table(rows=2, cols=2)
    fill(flagged, 0, "Fee", "Amount")
    fill(flagged, 1, "Setup", "1 percent")
    flagged.rows[0]._tr.get_or_add_trPr().append(parse_xml(f"<w:tblHeader {W}/>"))
    doc.add_paragraph("-")

    # Key-value table: the first row is data, not a header.
    pairs = doc.add_table(rows=2, cols=2)
    fill(pairs, 0, "Version", "1.2")
    fill(pairs, 1, "Effective date", "1 March 2025")
    doc.add_paragraph("-")

    # Merged title row, horizontal merge in the header, vertical merge in the body.
    merged = doc.add_table(rows=4, cols=3)
    merged.cell(0, 0).merge(merged.cell(0, 2)).text = "Fee schedule"
    fill(merged, 1, "Segment", "Item", "Fee", bold=True)
    fill(merged, 2, "Retail", "Setup", "100")
    fill(merged, 3, "", "Annual", "200")
    merged.cell(2, 0).merge(merged.cell(3, 0))
    doc.add_paragraph("-")

    # Table nested inside a cell.
    outer = doc.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "Documents"
    cell = outer.cell(0, 1)
    cell.paragraphs[0].text = "KYC"
    cell.add_table(rows=1, cols=2).cell(0, 0).text = "GST"

    text = docx_result(doc, tmp_path)["text"]

    assert text.split("\n\n-\n\n") == [
        "Fee: Setup | Amount: 1 percent",
        "Version | 1.2\nEffective date | 1 March 2025",
        "Fee schedule\nSegment: Retail | Item: Setup | Fee: 100\n"
        "Segment: Retail | Item: Annual | Fee: 200",
        "Documents | KYC GST",
    ]
    assert extract_metadata(text)[0] == "1.2"


def test_docx_tracked_changes_toc_content_controls_and_page_furniture(tmp_path):
    doc = Document()
    doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.sections[0].header.paragraphs[0].text = "CONFIDENTIAL"
    doc.sections[0].footer.paragraphs[0].text = "Page footer"
    doc.element.body.insert(0, parse_xml(
        f"<w:sdt {W}><w:sdtContent><w:p><w:r><w:t>In a content control</w:t></w:r></w:p>"
        f"</w:sdtContent></w:sdt>"))
    doc.add_paragraph("Eligibility ........ 3", style="TOC 1")

    rate = doc.add_paragraph("Rate is ")
    rate._p.append(parse_xml(
        f'<w:ins {W} w:id="1" w:author="a"><w:r><w:t>14 percent</w:t></w:r></w:ins>'))
    rate._p.append(parse_xml(
        f'<w:del {W} w:id="2" w:author="a"><w:r><w:delText>12 percent</w:delText></w:r></w:del>'))
    rate._p.append(parse_xml(
        f'<w:hyperlink {W} w:anchor="x"><w:r><w:t> (see policy)</w:t></w:r></w:hyperlink>'))

    lines = doc.add_paragraph("Line one")
    lines.add_run().add_break()
    lines.add_run("Line two")

    assert docx_result(doc, tmp_path)["text"] == (
        "In a content control\n\nRate is 14 percent (see policy)\n\nLine one\nLine two"
    )


def test_docx_heading_by_style_inheritance_and_outline_level(tmp_path):
    doc = Document()
    based = doc.styles.add_style("Policy Heading", WD_STYLE_TYPE.PARAGRAPH)
    based.base_style = doc.styles["Heading 2"]
    outline = doc.styles.add_style("Article", WD_STYLE_TYPE.PARAGRAPH)
    outline.element.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {W} w:val="2"/>'))
    unnumbered = doc.add_paragraph("Plain text")
    unnumbered._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {W}><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'))
    doc.add_paragraph("Clause one", style="Policy Heading")
    doc.add_paragraph("Clause two", style="Article")

    assert docx_result(doc, tmp_path)["text"] == "Plain text\n\n## Clause one\n### Clause two"


def test_docx_bad_legacy_and_empty_files_are_reported_not_fatal(tmp_path):
    (tmp_path / "bad.docx").write_bytes(b"\xd0\xcf\x11\xe0 an old .doc renamed")
    (tmp_path / "old.doc").write_bytes(b"x")
    Document().save(tmp_path / "empty.docx")

    results = {r["source"]: r for r in extract_directory(tmp_path)}

    assert results["bad.docx"]["status"] == "error"
    assert "not a valid .docx" in results["bad.docx"]["message"]
    assert results["old.doc"]["status"] == "skipped"
    assert "re-save as .docx" in results["old.doc"]["message"]
    assert results["empty.docx"]["status"] == "skipped"